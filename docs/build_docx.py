#!/usr/bin/env python
"""把 docs/ 下的 Markdown 技术方案编译成一份带目录的 Word 文档。

用法（在项目根目录执行）：

    python docs/build_docx.py

产出：

    OfferMind 技术方案.docx      —— 合并后的 Word 文档（图片已内嵌）
    docs/images/*.png            —— 22 张 Mermaid 图的高分辨率渲染（同时可供 PPT 复用）

依赖：

    pandoc（md → docx）、node/npx（mermaid-cli 渲染图）、python-docx、Pillow

注意：脚本会重建 docs/images/ 与临时目录 .docxbuild/。修改任何 docs/*.md 后重跑即可。
"""
from __future__ import annotations

import json
import pathlib
import re
import shutil
import subprocess
import sys

ROOT = pathlib.Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
BUILD = ROOT / ".docxbuild"
OUT = ROOT / "OfferMind 技术方案.docx"

# 章节顺序（README 作为导读放最前）
ORDER = ["README"] + [f"{i:02d}" for i in range(15)] + ["答辩PPT大纲"]

# 不并入 Word 文档的章节。源文件仍保留在 docs/ 下，只是不参与合并
# —— 答辩 PPT 大纲是答辩准备材料，不属于技术方案正文。
EXCLUDE_FROM_DOCX = {"答辩PPT大纲"}

# 并入时在该标题处截断（连同该标题及其后的全部内容一起丢弃）。
# 键是文件名（不含 .md）。附录第 7 节「相关文档索引」是指向其他文档的导航，
# 合并成单文档后没有意义。
TRUNCATE_AT = {"14-附录": "## 7. 相关文档索引"}

# emoji → Word 中更稳妥的符号（中文字体不一定覆盖彩色 emoji）
EMOJI = {"✅": "✓", "❌": "✗", "⚠️": "▲", "⚠": "▲", "⏳": "◇",
         "📋": "▪", "🔴": "●", "🟢": "●", "🟡": "●", "💡": "※"}

PAGEBREAK = '\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n'

TITLE_BLOCK = """---
title: "OfferMind 技术方案"
subtitle: "可观测 Agentic RAG 驱动的 AI 求职面试训练平台"
author: "马乾恩　蒋志敏　李贤星"
date: "大模型部署与微调 · 课程设计 + 竞赛项目"
lang: zh-CN
---

"""

TEXT_WIDTH_IN = 5.83    # Word 正文区宽度（A4 + 默认页边距）
MAX_HEIGHT_IN = 8.2     # 图片高度上限。留出图注与行距的余量，否则图会被推到下一页、
                        # 前一页只剩一个孤零零的图注（实测 8.6in 时会出现这种空页）


def resolve(exe: str) -> str:
    """解析出可执行文件的绝对路径。

    Windows 的 CreateProcess 不会自动补 PATHEXT 里的扩展名，所以像 npx 这种
    只有 .CMD 版本的工具，直接传裸名会报 WinError 2，必须先 which 一下。
    """
    return shutil.which(exe) or exe


def find_sources(include_excluded: bool = False) -> list[pathlib.Path]:
    """按 ORDER 找出 Markdown 文件。

    include_excluded=True 时连同 EXCLUDE_FROM_DOCX 里的章节一起返回 —— 渲染图时需要，
    因为 docs/images/ 要对全部章节保持最新，哪怕某章不并入 Word。
    """
    files = []
    for stem in ORDER:
        if stem in EXCLUDE_FROM_DOCX and not include_excluded:
            continue
        for p in sorted(DOCS.glob("*.md")):
            if p.stem == stem or p.stem.startswith(stem + "-"):
                files.append(p)
                break
    return files


def truncate(text: str, stem: str) -> str:
    """按 TRUNCATE_AT 把某份文档从指定标题处截断。"""
    marker = TRUNCATE_AT.get(stem)
    if marker and marker in text:
        text = text.split(marker, 1)[0].rstrip()
    return text


def extract_mermaid(files: list[pathlib.Path]) -> list[dict]:
    """把每份文档里的 mermaid 块提取成独立 .mmd，返回渲染映射。"""
    (BUILD / "mmd").mkdir(parents=True, exist_ok=True)
    mapping, n = [], 0
    for f in files:
        text = f.read_text(encoding="utf-8")
        for idx, block in enumerate(re.findall(r"```mermaid\n(.*?)```", text, re.S), 1):
            n += 1
            ascii_name = f"fig{n:02d}"
            (BUILD / "mmd" / f"{ascii_name}.mmd").write_text(block, encoding="utf-8")
            mapping.append({"ascii": ascii_name, "stem": f.stem, "idx": idx,
                            "png": f"{f.stem}-{idx}.png", "src": str(f)})
    (BUILD / "map.json").write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
    return mapping


def render_diagrams(mapping: list[dict]) -> None:
    """用 mermaid-cli 把每个 .mmd 渲染成 PNG（3 倍分辨率，白底）。

    增量渲染：源 Markdown 比目标 .png 旧就跳过，避免每次重建都白跑几分钟。
    （比的是源 .md 的时间戳，不是临时 .mmd —— 后者每次都是新写的，永远最新。）
    """
    img_dir = DOCS / "images"
    img_dir.mkdir(parents=True, exist_ok=True)
    skipped = 0
    for m in mapping:
        mmd = BUILD / "mmd" / f"{m['ascii']}.mmd"
        target = img_dir / m["png"]
        src_mtime = pathlib.Path(m["src"]).stat().st_mtime
        if target.exists() and target.stat().st_mtime >= src_mtime:
            skipped += 1
            continue
        cmd = [resolve("npx"), "-y", "@mermaid-js/mermaid-cli",
               "-i", str(mmd), "-o", str(target), "-b", "white", "-s", "3"]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
        status = "✓" if target.exists() else "✗"
        print(f"  {status} {m['png']}")
        if not target.exists():
            print(r.stderr[-500:], file=sys.stderr)
    if skipped:
        print(f"  （{skipped} 张图无变化，已跳过渲染）")


def make_reference_docx() -> pathlib.Path:
    """生成带中文字体的 pandoc 样式模板。"""
    import docx
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.enum.style import WD_STYLE_TYPE

    base = BUILD / "ref-base.docx"
    base.write_bytes(subprocess.run(
        [resolve("pandoc"), "--print-default-data-file", "reference.docx"],
        capture_output=True, timeout=120).stdout)
    d = docx.Document(str(base))
    CJK, CJK_H, MONO = "宋体", "微软雅黑", "Consolas"

    def set_font(style, latin, cjk, size=None, bold=None, color=None):
        f = style.font
        f.name = latin
        if size:
            f.size = Pt(size)
        if bold is not None:
            f.bold = bold
        if color:
            f.color.rgb = color
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = rPr.makeelement(qn("w:rFonts"), {})
            rPr.insert(0, rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            rFonts.set(qn(attr), cjk if attr == "w:eastAsia" else latin)

    byname = {s.name.strip().lower(): s for s in d.styles}
    cfg = {
        "normal": (CJK, CJK, 10.5, None, None),
        "body text": (CJK, CJK, 10.5, None, None),
        "first paragraph": (CJK, CJK, 10.5, None, None),
        "compact": (CJK, CJK, 10.0, None, None),
        "heading 1": (CJK_H, CJK_H, 18, True, RGBColor(0x1F, 0x3A, 0x5F)),
        "heading 2": (CJK_H, CJK_H, 15, True, RGBColor(0x1F, 0x3A, 0x5F)),
        "heading 3": (CJK_H, CJK_H, 12.5, True, RGBColor(0x2E, 0x5A, 0x88)),
        "heading 4": (CJK_H, CJK_H, 11, True, RGBColor(0x2E, 0x5A, 0x88)),
        "heading 5": (CJK_H, CJK_H, 10.5, True, None),
        "block text": (CJK, CJK, 10, None, RGBColor(0x44, 0x44, 0x44)),
        "table": (CJK, CJK, 9, None, None),
        "table caption": (CJK, CJK, 9, None, None),
        "caption": (CJK, CJK, 9, None, None),
        "image caption": (CJK, CJK, 9, None, None),
        "title": (CJK_H, CJK_H, 26, True, RGBColor(0x1F, 0x3A, 0x5F)),
        "verbatim char": (MONO, CJK, 9.5, None, None),
    }
    for name, c in cfg.items():
        if name in byname:
            set_font(byname[name], *c)
    for want, c in [("Source Code", (MONO, CJK, 8.5, None, None)),
                    ("Verbatim", (CJK, CJK, 9.5, None, None))]:
        if want.lower() not in byname:
            set_font(d.styles.add_style(want, WD_STYLE_TYPE.PARAGRAPH), *c)

    out = BUILD / "ref.docx"
    d.save(str(out))
    return out


def convert(text: str, stem: str, mapping: list[dict]) -> str:
    """把一份 Markdown 转成适合 Word 的 Markdown。"""
    idx = [0]

    def sub_mermaid(_m):
        idx[0] += 1
        return f"\n![{stem} 图 {idx[0]}](images/{stem}-{idx[0]}.png)\n"

    text = re.sub(r"```mermaid\n.*?```", sub_mermaid, text, flags=re.S)
    text = re.sub(r"^\*\*上一篇\*\*.*$", "", text, flags=re.M)      # 去掉导航行
    text = re.sub(r"(?<!!)\[([^\]]+)\]\([^)]*\)", r"\1", text)      # 链接保留文字
    for k, v in EMOJI.items():
        text = text.replace(k, v)
    return re.sub(r"\n{4,}", "\n\n\n", text).strip()


def cap_tall_images(md: str) -> str:
    """给高宽比过大的图加高度约束，避免内嵌后在 Word 里溢出页面。"""
    from PIL import Image

    counter = [0]
    need: dict[int, float] = {}

    def probe(m):
        counter[0] += 1
        path = DOCS / m.group(2)
        if not path.exists():
            return m.group(0)
        w, h = Image.open(path).size
        if TEXT_WIDTH_IN * h / w > MAX_HEIGHT_IN:
            need[counter[0]] = MAX_HEIGHT_IN
        return m.group(0)

    re.sub(r"!\[([^\]]*)\]\((images/[^)]+)\)", probe, md)
    counter[0] = 0

    def apply(m):
        counter[0] += 1
        if counter[0] in need:
            return f"{m.group(1)}{{height={MAX_HEIGHT_IN}in}}"
        return m.group(0)

    return re.sub(r"(!\[[^\]]*\]\(images/[^)]+\))", apply, md)


def export_pdf() -> pathlib.Path | None:
    """用 Word COM 把 docx 导出为 PDF。

    走 Word 而不是 LibreOffice，有两个好处：保真度最高（就是 Word 自己的排版），
    且能在导出前刷新目录域，让 PDF 里的目录带真实页码。同时按标题生成 PDF 书签。
    需要 Windows + 已安装 Word + pywin32。不可用时返回 None（不视为失败）。
    """
    pdf = OUT.with_suffix(".pdf")
    try:
        import win32com.client as win32
    except ImportError:
        print("  ⚠ 未安装 pywin32，跳过 PDF 导出（可 pip install pywin32 后重跑）")
        return None

    word = None
    try:
        word = win32.gencache.EnsureDispatch("Word.Application")
    except Exception:
        try:
            word = win32.Dispatch("Word.Application")
        except Exception as exc:
            print(f"  ⚠ 无法启动 Word，跳过 PDF 导出：{exc}")
            return None

    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(str(OUT), ReadOnly=False, AddToRecentFiles=False)
        for i in range(1, doc.TablesOfContents.Count + 1):
            doc.TablesOfContents(i).Update()
        doc.Fields.Update()
        doc.Repaginate()
        doc.Save()                       # 让 docx 里的目录页码也一并落盘
        doc.ExportAsFixedFormat(
            str(pdf), 17, False, 0, 0, 1, 1, 0,
            True, True, 1, True, True, False)   # 17=PDF, CreateBookmarks=1（按标题）
        doc.Close(SaveChanges=0)
    finally:
        word.Quit()
    return pdf if pdf.exists() else None


def main() -> int:
    global OUT
    # --to <路径>：改写到别的路径。用于目标文件被 Word/WPS 打开时先构建到临时位置，
    # 关闭后再移动过去，避免白等一次完整的 PDF 导出。
    if "--to" in sys.argv:
        OUT = pathlib.Path(sys.argv[sys.argv.index("--to") + 1]).resolve()

    if BUILD.exists():
        shutil.rmtree(BUILD)
    BUILD.mkdir(parents=True)

    all_files = find_sources(include_excluded=True)
    docx_files = [f for f in all_files if f.stem not in EXCLUDE_FROM_DOCX]
    if len(docx_files) != len(ORDER) - len(EXCLUDE_FROM_DOCX):
        print(f"⚠ 合并清单为 {len(docx_files)} 份，预期 "
              f"{len(ORDER) - len(EXCLUDE_FROM_DOCX)} 份，请检查 docs/ 下的文件名", file=sys.stderr)

    # 图仍然全量渲染（含被排除章节的图），保证 docs/images/ 始终是最新的
    print(f"[1/5] 提取 Mermaid 图（{len(all_files)} 份文档）")
    mapping = extract_mermaid(all_files)
    print(f"      共 {len(mapping)} 张")

    print("[2/5] 渲染图（mermaid-cli，3 倍分辨率，增量）")
    render_diagrams(mapping)

    print(f"[3/5] 合并 Markdown（{len(docx_files)} 份）")
    for f in docx_files:
        if f.stem in TRUNCATE_AT:
            print(f"      截断 {f.name} 于「{TRUNCATE_AT[f.stem]}」")
    parts = [convert(truncate(f.read_text(encoding="utf-8"), f.stem), f.stem, mapping)
             for f in docx_files]
    body = PAGEBREAK.join(parts)
    body = body.replace("# OfferMind 技术方案\n\n> **可观测", "# 方案导读\n\n> **可观测", 1)
    combined = TITLE_BLOCK + cap_tall_images(body) + "\n"
    (BUILD / "combined.md").write_text(combined, encoding="utf-8")
    print(f"      {len(combined.splitlines())} 行")

    if "--dry-run" in sys.argv:
        print("\n--dry-run：已停在合并步骤，未生成 Word。合并结果见 "
              f"{BUILD / 'combined.md'}")
        return 0

    print("[4/5] 生成 Word（pandoc + 中文样式模板）")
    ref = make_reference_docx()
    cmd = [resolve("pandoc"), str(BUILD / "combined.md"), "-o", str(OUT),
           f"--reference-doc={ref}", "--toc", "--toc-depth=3",
           f"--resource-path={DOCS}", "--from=markdown+raw_attribute+pipe_tables",
           "--standalone"]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    if r.returncode != 0:
        print(r.stderr, file=sys.stderr)
        return 1

    import docx
    from docx.shared import Emu
    d = docx.Document(str(OUT))
    tallest = max(Emu(s.height).inches for s in d.inline_shapes) if d.inline_shapes else 0
    print(f"\n✓ 已生成 {OUT.name}")
    print(f"  段落 {len(d.paragraphs)} / 表格 {len(d.tables)} / 图片 {len(d.inline_shapes)}")
    print(f"  最高图片 {tallest:.2f} in（上限 {MAX_HEIGHT_IN} in）")

    if "--no-pdf" not in sys.argv:
        print("[5/5] 导出 PDF（Word COM，会刷新目录页码，约需数分钟）")
        pdf = export_pdf()
        if pdf:
            print(f"\n✓ 已生成 {pdf.name}（{pdf.stat().st_size / 1024 / 1024:.2f} MB，"
                  f"含按标题生成的 PDF 书签）")
    else:
        print("[5/5] 已按要求跳过 PDF 导出（--no-pdf）")

    shutil.rmtree(BUILD, ignore_errors=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
